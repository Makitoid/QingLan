"""0.4.1 F3：教师名单的批量导入（xlsx / txt / csv）与导出（xlsx / csv）。

契约以代码实际行为为准（app/api/admin.py 的 ``process_teacher_import_rows`` /
``export_teachers``、app/services/export.py 的 ``TEACHERS_EXPORT_*``）：
- 导入列：**工号 | 姓名 | 可教组别**（3 列，组别列可选）。解析与学生导入共用
  ``parse_import_file``，只是列数与表头首格白名单不同；多组别分隔符同组别列，
  不存在的组自动创建，生效方式是幂等写 ``teacher_groups``（层 2）。
- 失败行 reason：「字段缺失」「工号重复」（学号/工号在同一张 users 表里全局唯一）。
- 建号密码对齐 PW-01：整批统一 ``config.DEFAULT_INITIAL_PASSWORD`` + 首登强制改密。
- 审计（AU-04）：逐人一条 ``teacher_create_pw`` + 整批一条 ``teacher_import``。
- 导出列：``工号 | 姓名 | 可教组别 | 名单学生数 | 状态``，与学生导出同一套
  Content-Disposition（ASCII 兜底 + RFC5987）；``q`` 与列表接口同口径。
"""
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

import app.core.config as config
from app.core.security import create_token
from app.main import app
from app.models import (AuditLog, Group, GroupMember, TeacherGroup, User)
from app.services.export import (TEACHERS_EXPORT_HEADER, build_teachers_csv,
                                 build_teachers_xlsx)

INITIAL = config.DEFAULT_INITIAL_PASSWORD
IMPORT_URL = "/api/admin/teachers/import"
EXPORT_URL = "/api/admin/teachers/export"


def _add(db, obj):
    db.add(obj)
    db.commit()
    db.refresh(obj)
    return obj


def _upload(client, headers, filename, content, ctype="text/plain"):
    return client.post(IMPORT_URL, headers=headers, files={"file": (filename, content, ctype)})


def _xlsx_bytes(rows):
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _teachers(db):
    return {u.username: u for u in db.query(User).filter(User.role == "teacher").all()}


def _group_names(db):
    return sorted(g.name for g in db.query(Group).all())


def _teacher_group_pairs(db):
    db.expire_all()
    rows = db.query(TeacherGroup).all()
    by_id = {u.id: u.username for u in db.query(User).all()}
    groups = {g.id: g.name for g in db.query(Group).all()}
    return sorted((by_id[r.teacher_id], groups[r.group_id]) for r in rows)


def _audit(db, action):
    db.expire_all()
    return db.query(AuditLog).filter(AuditLog.action == action).all()


@pytest.fixture(autouse=True)
def fast_hash(monkeypatch):
    """bcrypt 太慢；替身让「整批共用同一份初始密码哈希」可被断言。"""
    monkeypatch.setattr("app.api.admin.hash_password", lambda pwd: f"pw<{pwd}>")


@pytest.fixture()
def client():
    return TestClient(app)


@pytest.fixture()
def admin(db):
    return _add(db, User(username="root", password_hash="x", role="admin",
                         display_name="管理员", must_change_password=0))


@pytest.fixture()
def admin_headers(admin):
    return {"Authorization": f"Bearer {create_token(admin)}"}


# ================================================ 导入

class TestTeacherImport:
    def test_pipe_separated_with_groups_creates_accounts_and_layer2(self, db, client, admin_headers):
        resp = _upload(client, admin_headers, "t.txt",
                       "T01|王老师|甲组\nT02|李老师|甲组、乙组\n".encode("utf-8"))
        assert resp.status_code == 200, resp.text
        assert resp.json() == {"success_count": 2, "failures": []}

        teachers = _teachers(db)
        assert sorted(teachers) == ["T01", "T02"]
        assert teachers["T01"].display_name == "王老师"
        # PW-01：统一初始密码 + 首登强制改密
        assert {t.password_hash for t in teachers.values()} == {f"pw<{INITIAL}>"}
        assert all(t.must_change_password == 1 for t in teachers.values())
        assert _group_names(db) == ["乙组", "甲组"]
        assert _teacher_group_pairs(db) == [("T01", "甲组"), ("T02", "乙组"), ("T02", "甲组")]

    def test_group_column_is_optional(self, db, client, admin_headers):
        resp = _upload(client, admin_headers, "t.txt", "T09|赵老师\n".encode("utf-8"))
        assert resp.json() == {"success_count": 1, "failures": []}
        assert _group_names(db) == []
        assert _teacher_group_pairs(db) == []

    def test_header_row_is_skipped(self, db, client, admin_headers):
        resp = _upload(client, admin_headers, "t.csv",
                       "工号,姓名,可教组别\nT11,陈老师,丙组\n".encode("utf-8"))
        assert resp.json() == {"success_count": 1, "failures": []}
        assert sorted(_teachers(db)) == ["T11"]

    def test_duplicate_username_fails_whole_row(self, db, client, admin_headers, teacher):
        resp = _upload(client, admin_headers, "t.txt",
                       f"{teacher.username}|重名|\nT12|新用户|\n".encode("utf-8"))
        body = resp.json()
        assert body["success_count"] == 1
        assert [(f["line"], f["username"], f["reason"]) for f in body["failures"]] == [
            (1, teacher.username, "工号重复")]
        assert sorted(_teachers(db)) == ["T12", teacher.username]

    def test_missing_columns_are_reported_per_row(self, db, client, admin_headers):
        resp = _upload(client, admin_headers, "t.txt", "T13||甲组\n".encode("utf-8"))
        body = resp.json()
        assert body["success_count"] == 0
        assert [(f["line"], f["reason"]) for f in body["failures"]] == [(1, "字段缺失")]
        # 失败行的组别不产生副作用
        assert _group_names(db) == []

    def test_xlsx_import(self, db, client, admin_headers):
        resp = client.post(IMPORT_URL, headers=admin_headers, files={
            "file": ("t.xlsx", _xlsx_bytes([["工号", "姓名", "可教组别"],
                                            ["T21", "周老师", "甲组"],
                                            ["T22", "吴老师", None]]),
                     "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert resp.status_code == 200, resp.text
        assert resp.json() == {"success_count": 2, "failures": []}
        assert sorted(_teachers(db)) == ["T21", "T22"]

    def test_extra_columns_merge_into_group_column(self, db, client, admin_headers):
        # 教师模板只有 3 列：第 4 列及之后并入组别列，不丢数据（与学生侧同规则）
        resp = _upload(client, admin_headers, "t.txt", "T31|郑老师|甲组|备注\n".encode("utf-8"))
        assert resp.json() == {"success_count": 1, "failures": []}
        assert _teacher_group_pairs(db) == [("T31", "备注"), ("T31", "甲组")]

    def test_file_validation_reuses_student_rules(self, client, admin_headers):
        assert _upload(client, admin_headers, "t.docx", b"T1|a").status_code == 415
        assert _upload(client, admin_headers, "t.txt", b"").status_code == 422
        too_big = _upload(client, admin_headers, "t.txt", b"T1|a\n" * 2_000_000)
        assert too_big.status_code == 413

    def test_audit_trail(self, db, client, admin_headers):
        resp = _upload(client, admin_headers, "t.txt",
                       "T41|甲老师|甲组\nT42|乙老师|\nT43||\n".encode("utf-8"))
        assert resp.json()["success_count"] == 2
        assert len(_audit(db, "teacher_create_pw")) == 2
        batch = _audit(db, "teacher_import")
        assert len(batch) == 1
        assert '"success_count": 2' in batch[0].detail
        assert '"failure_count": 1' in batch[0].detail

    def test_requires_admin(self, client, db, teacher):
        headers = {"Authorization": f"Bearer {create_token(teacher)}"}
        assert _upload(client, headers, "t.txt", b"T1|a").status_code == 403
        assert _upload(client, None, "t.txt", b"T1|a").status_code == 401


# ================================================ 导出

class TestTeacherExport:
    def _seed(self, db, admin, teacher, student):
        group = _add(db, Group(name="甲组"))
        db.add(GroupMember(group_id=group.id, student_id=student.id))
        db.add(TeacherGroup(teacher_id=teacher.id, group_id=group.id))
        _add(db, User(username="t002", password_hash="x", role="teacher",
                      display_name="李老师", is_active=0, must_change_password=0))
        db.commit()

    def test_xlsx_columns_and_rows(self, db, client, admin_headers, admin, teacher, student):
        self._seed(db, admin, teacher, student)
        resp = client.get(EXPORT_URL, headers=admin_headers)
        assert resp.status_code == 200, resp.text
        assert resp.headers["content-type"] == \
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert "filename*=UTF-8''" in resp.headers["content-disposition"]

        rows = list(load_workbook(io.BytesIO(resp.content)).active.iter_rows(values_only=True))
        assert list(rows[0]) == TEACHERS_EXPORT_HEADER
        assert list(rows[1]) == ["t001", "王老师", "甲组", 1, "启用"]
        assert rows[2][0:2] == ("t002", "李老师") and rows[2][4] == "已停用"

    def test_csv_has_bom_and_same_cells(self, db, client, admin_headers, admin, teacher, student):
        self._seed(db, admin, teacher, student)
        resp = client.get(EXPORT_URL, headers=admin_headers, params={"format": "csv"})
        assert resp.status_code == 200, resp.text
        text = resp.content.decode("utf-8-sig")
        assert text.splitlines()[0].startswith("工号,姓名,可教组别")
        assert "t001,王老师,甲组,1,启用" in text

    def test_q_filter_matches_list_endpoint(self, db, client, admin_headers, admin, teacher, student):
        self._seed(db, admin, teacher, student)
        rows = load_workbook(io.BytesIO(
            client.get(EXPORT_URL, headers=admin_headers, params={"q": "李"}).content
        )).active
        cells = list(rows.iter_rows(values_only=True))
        assert len(cells) == 2 and cells[1][0] == "t002"

    def test_bad_format_is_422(self, client, db, admin_headers):
        assert client.get(EXPORT_URL, headers=admin_headers,
                          params={"format": "pdf"}).status_code == 422

    def test_requires_admin(self, client, db, teacher):
        headers = {"Authorization": f"Bearer {create_token(teacher)}"}
        assert client.get(EXPORT_URL, headers=headers).status_code == 403
        assert client.get(EXPORT_URL).status_code == 401

    def test_builders_are_pure(self, db):
        rows = [{"username": "t1", "name": "甲", "groups": ["A", "B"],
                 "student_count": 3, "is_active": True}]
        assert build_teachers_csv(rows).decode("utf-8-sig").splitlines()[1] == "t1,甲,A、B,3,启用"
        assert len(build_teachers_xlsx(rows)) > 0
