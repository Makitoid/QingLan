"""成绩导出：openpyxl 生成 xlsx。

导出列随题单动态展开（SC-02）：固定「学号/姓名/提交次数/最高单题分/总分」+
每题得分一列（列头=题目标题，顺序=题单 seq）+ 末尾「最后提交时间」。
openpyxl 采用函数内延迟 import —— 未安装该依赖时后端进程仍可启动，
只有真正调用导出/解析接口才报错，避免拖垮整个服务。
"""
from __future__ import annotations

from datetime import datetime, timedelta
from urllib.parse import quote

from ..core.security import APIError

XLSX_SHEET_MAX_LEN = 31
XLSX_SHEET_ILLEGAL = "[]:*?/\\"
# SC-02：固定列（前 5 列）+ 每题得分动态列 + 末尾「最后提交时间」列
XLSX_FIXED_HEADER = ["学号", "姓名", "提交次数", "最高单题分", "总分"]
XLSX_TIME_HEADER = "最后提交时间"
XLSX_FIXED_WIDTHS = [16, 16, 12, 14, 14]
XLSX_PROBLEM_WIDTH = 12
XLSX_TIME_WIDTH = 22
TIME_FMT = "%Y-%m-%d %H:%M:%S"

MISSING_OPENPYXL = APIError(500, "OPENPYXL_MISSING", "服务器未安装 openpyxl，无法处理 Excel 文件")


def load_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        raise MISSING_OPENPYXL
    return openpyxl


def sheet_title(title: str) -> str:
    """Excel 限制：sheet 名 ≤31 字符且不含 []:*?/\\ ；空名回退为「成绩」。"""
    cleaned = "".join(ch for ch in (title or "") if ch not in XLSX_SHEET_ILLEGAL).strip()
    cleaned = cleaned[:XLSX_SHEET_MAX_LEN]
    return cleaned or "成绩"


def to_local(utc_text: str | None, tz_offset: int) -> str:
    """库内时间存 UTC（本项目已知坑 9）；导出按前端传入的分钟偏移换算本地时间。"""
    if not utc_text:
        return ""
    try:
        moment = datetime.strptime(utc_text[:19], TIME_FMT)
    except ValueError:
        return utc_text
    return (moment + timedelta(minutes=tz_offset)).strftime(TIME_FMT)


def _problem_columns(rows: list[dict]) -> list[dict]:
    """每题得分列的规格（problem_id + 列头标题），取自首条有题分明细的行。

    所有行的 problem_scores 都由 stats.student_rows 按同一题单（seq 顺序）生成，
    因此以其中一份为列顺序即可；rows 为空（无学生）时退化为不带每题列。
    """
    for row in rows:
        scores = row.get("problem_scores") or []
        if scores:
            return [
                {"problem_id": ps.get("problem_id"), "title": ps.get("title", "")}
                for ps in scores
            ]
    return []


def build_assignment_students_xlsx(assignment_title: str, rows: list[dict], tz_offset: int = 0) -> bytes:
    """场次学生成绩 → xlsx 字节串。

    列布局（SC-02）：`学号 | 姓名 | 提交次数 | 最高单题分 | 总分 | <每题得分列…> | 最后提交时间`。
    每题一列，列头=题目标题，列顺序即题单 seq；rows 为 services.stats.student_rows 的输出
    （含 username / name / submitted_count / best_effective_score / total_score /
    problem_scores / last_submitted_at）。
    """
    openpyxl = load_openpyxl()
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title(assignment_title)

    columns = _problem_columns(rows)
    ws.append([*XLSX_FIXED_HEADER, *(col["title"] for col in columns), XLSX_TIME_HEADER])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    widths = [
        *XLSX_FIXED_WIDTHS,
        *([XLSX_PROBLEM_WIDTH] * len(columns)),
        XLSX_TIME_WIDTH,
    ]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    for row in rows:
        by_problem = {
            ps.get("problem_id"): ps.get("effective_score")
            for ps in (row.get("problem_scores") or [])
        }
        cells = [
            row.get("username", ""),
            row.get("name", ""),
            row.get("submitted_count", 0),
            row.get("best_effective_score", 0),
            row.get("total_score", 0),
        ]
        # 未提交的题写空串：Excel 显示为空（openpyxl 回读为 None，见坑 23）
        cells.extend(
            "" if by_problem.get(col["problem_id"]) is None else by_problem[col["problem_id"]]
            for col in columns
        )
        cells.append(to_local(row.get("last_submitted_at"), tz_offset))
        ws.append(cells)

    from io import BytesIO
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def attachment_headers(filename: str) -> dict[str, str]:
    """Content-Disposition：RFC5987 filename* 传中文，另给 ASCII 兜底文件名。"""
    safe = "".join(ch for ch in filename if ch.isascii() and (ch.isalnum() or ch in "._-")) or "export"
    fallback = safe if safe.lower().endswith(".xlsx") else f"{safe}.xlsx"
    quoted = quote(filename if filename.lower().endswith(".xlsx") else f"{filename}.xlsx", safe="")
    return {"Content-Disposition": f"attachment; filename=\"{fallback}\"; filename*=UTF-8''{quoted}"}
