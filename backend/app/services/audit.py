"""统一审计入口（AU-03）。

业务代码一律经 `log_audit()` 写日志，不得直接 insert `audit_log`。
只 add 不 commit：与触发它的写操作同属一个事务，因此必须在业务 `db.commit()` 之前调用。
日志只追加，不提供任何 UPDATE/DELETE 路径（保留期清理是唯一的删除路径，见 `prune_expired`）。

总开关与保留期读自 `site_settings` 单行（id=1，0.3.2 F5）：
- `audit_enabled = 0` 时 `log_audit` 直接返回；判定结果进程内缓存，写设置时显式失效。
- `audit_retention_days` 非空时按天数清理旧行；NULL = 永久保存。

动作 / 对象类型的中文名只在 `AUDIT_ACTION_LABELS` / `AUDIT_TARGET_LABELS` 维护一份，
接口出参与 Excel 导出共用，避免中英两份字典漂移。
"""
from __future__ import annotations

import json
import time
from datetime import datetime, timedelta, timezone
from io import BytesIO
from typing import Any

from sqlalchemy import delete
from sqlalchemy.orm import Session

from . import export as export_svc
from ..models import AuditLog, SiteSetting, User

TIME_FMT = "%Y-%m-%d %H:%M:%S"

AUDIT_RETENTION_MIN_DAYS = 1
AUDIT_RETENTION_MAX_DAYS = 3650
PRUNE_INTERVAL_SECONDS = 3600
AUDIT_EXPORT_MAX_ROWS = 50000

AUDIT_ACTION_LABELS: dict[str, str] = {
    "student_create_pw": "新建学生（发放初始密码）",
    "student_reset_pw": "重置学生密码",
    "student_batch_reset_pw": "批量重置学生密码",
    "teacher_create_pw": "新建教师（发放初始密码）",
    "teacher_reset_pw": "重置教师密码",
    "user_change_password": "用户修改密码",
    "user_is_active_change": "账号停用/启用",
    "group_create": "新建分组",
    "group_update": "分组改名",
    "group_delete": "删除分组",
    "group_member_change": "组成员变更",
    "teacher_group_assign": "教师可教组别分配",
    "teacher_student_bind": "教师拉入学生",
    "teacher_student_unbind": "教师移出学生",
    "subgroup_create": "新建子分组",
    "subgroup_update": "子分组改名",
    "subgroup_delete": "删除子分组",
    "subgroup_member_change": "子分组名单变更",
    "student_import": "批量导入学生",
    "score_manual_adjust": "成绩手动调分",
}

AUDIT_TARGET_LABELS: dict[str, str] = {
    "user": "账号",
    "student": "学生",
    "teacher": "教师",
    "group": "分组",
    "subgroup": "子分组",
    "submission": "提交",
    "problem": "题目",
    "assignment": "场次",
}

AUDIT_EXPORT_HEADER = ["时间", "操作人", "工号", "动作", "对象类型", "对象 ID", "详情 JSON"]
AUDIT_EXPORT_WIDTHS = [22, 18, 16, 28, 12, 10, 60]
AUDIT_EXPORT_SHEET = "审计日志"

_enabled_cache: bool | None = None
_last_prune_monotonic: float | None = None


def invalidate_audit_cache() -> None:
    """`PUT /admin/settings` 写完开关/保留期后调用：开关重新读库，并重臂惰性清理守卫。"""
    global _enabled_cache, _last_prune_monotonic
    _enabled_cache = None
    _last_prune_monotonic = None


def audit_enabled(db: Session) -> bool:
    global _enabled_cache
    if _enabled_cache is None:
        row = db.get(SiteSetting, 1)
        _enabled_cache = True if row is None else bool(row.audit_enabled)
    return _enabled_cache


def action_label(action: str) -> str:
    return AUDIT_ACTION_LABELS.get(action, action)


def target_label(target_type: str) -> str:
    return AUDIT_TARGET_LABELS.get(target_type, target_type)


def retention_cutoff(db: Session) -> str | None:
    """保留期起点（UTC 串）；NULL / 未设置 = 无期限，返回 None。"""
    row = db.get(SiteSetting, 1)
    days = row.audit_retention_days if row is not None else None
    if not days:
        return None
    return (datetime.now(timezone.utc) - timedelta(days=int(days))).strftime(TIME_FMT)


def prune_expired(db: Session) -> int:
    """删除超出保留期的审计行，返回删除条数。不 commit，随调用方的事务一起落库。"""
    cutoff = retention_cutoff(db)
    if cutoff is None:
        return 0
    result = db.execute(delete(AuditLog).where(AuditLog.created_at < cutoff))
    return int(result.rowcount or 0)


def _maybe_prune(db: Session) -> None:
    global _last_prune_monotonic
    now = time.monotonic()
    if _last_prune_monotonic is not None and now - _last_prune_monotonic < PRUNE_INTERVAL_SECONDS:
        return
    _last_prune_monotonic = now
    prune_expired(db)


def log_audit(db: Session, actor: User | int | None, action: str, target_type: str,
              target_id: int | None = None, detail: dict[str, Any] | None = None) -> None:
    if not audit_enabled(db):
        return
    _maybe_prune(db)
    actor_id = actor.id if isinstance(actor, User) else actor
    db.add(AuditLog(
        actor_id=actor_id,
        action=action,
        target_type=target_type,
        target_id=target_id,
        detail=json.dumps(detail, ensure_ascii=False) if detail is not None else None,
    ))


def build_audit_logs_xlsx(rows: list[dict[str, Any]]) -> bytes:
    """审计行（dict）→ xlsx 字节串；列见 `AUDIT_EXPORT_HEADER`。"""
    openpyxl = export_svc.load_openpyxl()
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = AUDIT_EXPORT_SHEET
    ws.append(AUDIT_EXPORT_HEADER)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for idx, width in enumerate(AUDIT_EXPORT_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    for row in rows:
        ws.append([
            row.get("created_at") or "",
            row.get("actor_name") or "",
            row.get("actor_username") or "",
            action_label(row.get("action") or ""),
            target_label(row.get("target_type") or ""),
            "" if row.get("target_id") is None else row["target_id"],
            row.get("detail") or "",
        ])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
